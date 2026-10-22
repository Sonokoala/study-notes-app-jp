"""ダークテーマの 16:9 プレゼンを作る。ノート欄に読み上げ用の原稿を入れる。

仕様:
  {"title": "...", "subtitle": "...", "lang": "en" | "ja",
   "slides": [{"title": "...", "bullets": ["..."], "script": ["一息で読める1文", ...]}]}

箇条書きと原稿の中の **文字** はキーワード。
  英語 (en) -> 原稿では太字（発音練習アプリでは太字＝強く読む語として表示される）
  日本語 (ja) -> 原稿では下線、箇条書きでは太字
このアプリのプレゼンは英語 (en) で作る。
"""
import re
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

FONT = {"ja": "Meiryo UI", "en": "Segoe UI"}
BG, PANEL, TEXT, MUTED, ACCENT, ACCENT2 = "0F1B2D", "1B2E45", "E6EDF5", "9AA8B8", "6FB1FF", "FF7FA6"
W, H = 960.0, 540.0
rgb = RGBColor.from_string


def runs(s):
    out, on = [], False
    for part in re.split(r"(\*\*)", s):
        if part == "**":
            on = not on
        elif part:
            out.append((part, on))
    return out


def add_run(p, text, size, color, font, bold=False, underline=False):
    r = p.add_run()
    r.text = text
    f = r.font
    f.name, f.size, f.bold, f.underline = font, Pt(size), bold, underline
    if color:
        f.color.rgb = rgb(color)
    rpr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", font)
    return r


def box(slide, x, y, w, h):
    tf = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h)).text_frame
    tf.word_wrap = True
    return tf


def background(slide, color=BG):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(color)


def script_notes(slide, lines, lang, font):
    tf = slide.notes_slide.notes_text_frame
    tf.text = ""
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        for t, key in runs(line):
            if lang == "ja":
                add_run(p, t, 14, None, font, underline=key)
            else:
                add_run(p, t, 14, None, font, bold=key)


def build(spec, out):
    lang = spec.get("lang", "en")
    font = FONT.get(lang, FONT["en"])
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(W), Pt(H)
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    background(s, "0B1424")
    add_run(box(s, 60, 170, 840, 90).paragraphs[0], spec.get("title", ""), 40, "FFFFFF", font, bold=True)
    if spec.get("subtitle"):
        add_run(box(s, 60, 270, 840, 40).paragraphs[0], spec["subtitle"], 18, ACCENT, font)
    if spec.get("intro_script"):
        script_notes(s, spec["intro_script"], lang, font)

    for i, sl in enumerate(spec.get("slides", []), 1):
        s = prs.slides.add_slide(blank)
        background(s)
        bar = s.shapes.add_shape(1, Pt(0), Pt(0), Pt(8), Pt(H))
        bar.fill.solid(); bar.fill.fore_color.rgb = rgb(ACCENT); bar.line.fill.background()
        add_run(box(s, 48, 30, 864, 60).paragraphs[0], sl.get("title", ""), 30, "FFFFFF", font, bold=True)
        panel = s.shapes.add_shape(1, Pt(48), Pt(104), Pt(864), Pt(380))
        panel.fill.solid(); panel.fill.fore_color.rgb = rgb(PANEL); panel.line.fill.background()
        tf = box(s, 72, 124, 816, 340)
        bullets = sl.get("bullets", [])
        size = 24 if len(bullets) <= 4 else 20
        for j, b in enumerate(bullets):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.space_after = Pt(12)
            add_run(p, "▸ ", size, ACCENT, font)
            for t, key in runs(b):
                add_run(p, t, size, ACCENT2 if key else TEXT, font, bold=key)
        add_run(box(s, 840, 500, 80, 24).paragraphs[0], str(i), 11, MUTED, font)
        script_notes(s, sl.get("script", []), lang, font)

    prs.core_properties.title = spec.get("title", "")
    prs.save(out)
    return out
